from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class Document:
    source: str
    text: str


def load_file(path: Path) -> Document:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8")
    elif suffix == ".pdf":
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("读取 PDF 需要安装 pymupdf") from exc
        with fitz.open(path) as pdf:
            text = "\n".join(page.get_text() for page in pdf)
    elif suffix == ".docx":
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError("读取 DOCX 需要安装 python-docx") from exc
        doc = DocxDocument(path)
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    else:
        raise ValueError(f"不支持的文件类型: {path.name}")
    return Document(source=path.name, text=clean_text(text))


def load_directory(directory: Path) -> list[Document]:
    documents: list[Document] = []
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in {".txt", ".md", ".pdf", ".docx"}:
            documents.append(load_file(path))
    return documents


def clean_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()
